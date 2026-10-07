import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def create_element(name):
    return OxmlElement(name)

def add_page_number(run):
    fldChar1 = create_element('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = create_element('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = "PAGE"
    fldChar2 = create_element('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = create_element('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

doc = docx.Document()

# Page Margins - 0.75 in
sections = doc.sections
for section in sections:
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

# Normal Style
style = doc.styles['Normal']
font = style.font
font.name = 'Arial'
font.size = Pt(10.5)
font.color.rgb = RGBColor(0x11, 0x18, 0x27)

# --- Header Title ---
p_top = doc.add_paragraph()
p_top.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_top = p_top.add_run("Social Media Analytics Lab")
run_top.font.bold = True
run_top.font.size = Pt(14)
run_top.font.name = 'Arial'

p_top.paragraph_format.space_after = Pt(14)

# --- Student Details Table ---
table = doc.add_table(rows=4, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False

# Set Column Widths
col_widths = [Inches(3.4), Inches(3.4)]
for row in table.rows:
    for i, w in enumerate(col_widths):
        row.cells[i].width = w

details = [
    ("Name: Chinmay S. Santosh", "DOP: 06/10/2026"),
    ("Roll No: 711", "DOS: 10/10/2026"),
    ("Division: B", "Score:"),
    ("Group: 6", "")
]

for row_idx, (left_text, right_text) in enumerate(details):
    # Left Cell
    cell_l = table.cell(row_idx, 0)
    p_l = cell_l.paragraphs[0]
    p_l.paragraph_format.space_after = Pt(2)
    p_l.paragraph_format.space_before = Pt(2)
    r_l = p_l.add_run(left_text)
    r_l.font.bold = True
    r_l.font.size = Pt(11)

    # Right Cell
    cell_r = table.cell(row_idx, 1)
    p_r = cell_r.paragraphs[0]
    p_r.paragraph_format.space_after = Pt(2)
    p_r.paragraph_format.space_before = Pt(2)
    r_r = p_r.add_run(right_text)
    r_r.font.bold = True
    r_r.font.size = Pt(11)

doc.add_paragraph().paragraph_format.space_after = Pt(4)

# --- Horizontal Separator Line ---
p_hr = doc.add_paragraph()
p_hr.paragraph_format.space_before = Pt(4)
p_hr.paragraph_format.space_after = Pt(12)
pBdr = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                 r'<w:bottom w:val="single" w:sz="12" w:space="1" w:color="111827"/>'
                 r'</w:pBdr>')
p_hr._p.get_or_add_pPr().append(pBdr)

# --- Section 1: Project Title ---
p_sec1 = doc.add_paragraph()
r_sec1 = p_sec1.add_run("Project Title")
r_sec1.font.bold = True
r_sec1.font.size = Pt(12)
p_sec1.paragraph_format.space_after = Pt(2)

p_title = doc.add_paragraph()
r_title = p_title.add_run("Catalog Maker: High-Speed Product Catalog Platform")
r_title.font.bold = True
r_title.font.size = Pt(11)
r_title.font.color.rgb = RGBColor(0x1D, 0x4E, 0xD8)
p_title.paragraph_format.space_after = Pt(10)

# --- Section 2: Approach ---
p_sec2 = doc.add_paragraph()
r_sec2 = p_sec2.add_run("Approach")
r_sec2.font.bold = True
r_sec2.font.size = Pt(12)
p_sec2.paragraph_format.space_after = Pt(4)

app_text = (
    "The Catalog Maker platform is engineered as a reusable, mobile-first product catalog system designed for ultra-fast "
    "browsing speed and instant conversion. Instead of issuing live, blocking requests to external e-commerce platforms during customer pageviews, "
    "the architecture decouples the frontend from external systems using a central, lightweight internal database and REST API.\n\n"
    "Key Methodology & Technical Steps:\n"
    "1. Multi-Source E-Commerce Import & Normalization: Built dedicated API importers for Shopify (55 items), WooCommerce (55 items), "
    "and IndiaMART B2B (25 items), yielding a combined repository of 135 products normalized into a unified schema (SKUs, price, stock status, images, variants, and metadata).\n"
    "2. High-Speed Internal Database & REST API: Implemented a persistent file-backed JSON/SQLite database layer offering sub-5ms query response times. "
    "Exposed lightweight endpoints (/api/v1/products, /api/v1/categories, /api/v1/search, /api/v1/config) that decouple frontend rendering from Shopify/WooCommerce response schemas.\n"
    "3. Catalog Synchronization & Source Mutation Simulator: Built an asynchronous sync engine that detects price drops, stock updates, and new items "
    "while recording audit history logs. Developed a source mutation tool to demonstrate real-time external store updates propagating to the catalog.\n"
    "4. Dual Catalog Design Engine: Engineered two dynamically switchable frontend themes without data rebuilding: 'Rajdhani Royal Luxury' "
    "(Obsidian emerald & gold dark theme inspired by rajdhanicarpets.com) and 'Velox Glassmorphic' (acrylic glass light theme).\n"
    "5. WhatsApp Enquiry Engine & Wishlist: Designed an account-free wishlist (localStorage) and a multi-product selection bar that formats a pre-filled "
    "WhatsApp message listing selected titles, SKUs, pricing, and links targeting +919833113449.\n"
    "6. Deployment & Hosting: Prepared Express serverless API adapters (api/index.js) and vercel.json configurations for seamless Vercel deployment and GitHub integration."
)

p_app = doc.add_paragraph()
r_app = p_app.add_run(app_text)
r_app.font.size = Pt(9.5)
p_app.paragraph_format.space_after = Pt(12)

# --- Section 3: Screenshots ---
p_sec3 = doc.add_paragraph()
r_sec3 = p_sec3.add_run("Screenshots")
r_sec3.font.bold = True
r_sec3.font.size = Pt(12)
p_sec3.paragraph_format.space_after = Pt(6)

screenshots_data = [
    ("Figure 1: Rajdhani Royal Luxury Public Catalog Interface", 
     "Features obsidian gold aesthetic, instant search input, touch-scrollable category navigation, live in-stock badges, and sub-5ms API response latency."),
    
    ("Figure 2: Product Pop-Up Modal & High-Resolution Image Lightbox", 
     "Displays product specifications, variant choices, related recommendations, fullscreen image zoom, and direct WhatsApp inquiry trigger."),
    
    ("Figure 3: Multi-Select WhatsApp Enquiry Floating Action Bar", 
     "Enables multi-item selection with estimated subtotal calculations, generating a pre-filled WhatsApp inquiry targeting +919833113449."),
    
    ("Figure 4: Administrator Management Portal & Sync Engine Control", 
     "Overview metrics dashboard, product CRUD, external store data mutation simulator, and live sync execution audit logs.")
]

for title, desc in screenshots_data:
    p_box = doc.add_paragraph()
    p_box.paragraph_format.space_before = Pt(4)
    p_box.paragraph_format.space_after = Pt(4)
    
    r_fig = p_box.add_run(f"• {title}\n")
    r_fig.font.bold = True
    r_fig.font.size = Pt(9.5)
    r_fig.font.color.rgb = RGBColor(0x1D, 0x4E, 0xD8)
    
    r_desc = p_box.add_run(f"   {desc}")
    r_desc.font.size = Pt(9)
    r_desc.font.italic = True
    r_desc.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)

# Save Document
doc.save("Catalog_Maker_Submission.docx")
print("SUCCESS: Catalog_Maker_Submission.docx generated!")
